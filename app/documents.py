from copy import deepcopy
from pathlib import Path
from datetime import date
from decimal import Decimal
import tempfile, subprocess, os, io
from urllib.request import urlopen
from urllib.parse import urlparse, unquote
from PIL import Image
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx import Document
from docx.table import _Row

BASE = Path(__file__).resolve().parent
MAX_LOGO_BYTES = 10 * 1024 * 1024

def text(p, value):
    # Retain the paragraph and first run's formatting.
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]: run.text = ''
    else: p.add_run(value)

def cell(c, value):
    text(c.paragraphs[0], value)
    for p in c.paragraphs[1:]: text(p, '')

def uk(value): return date.fromisoformat(value).strftime('%d/%m/%Y')
def money(pence): return f'£{Decimal(pence)/100:,.2f}'


def company_logo(document):
    source = os.getenv('COMPANY_LOGO', '').strip()
    if not source:
        return  # Keep the template's existing logo.
    try:
        parsed = urlparse(source)
        if parsed.scheme in ('http', 'https'):
            with urlopen(source, timeout=10) as response:
                raw = response.read(MAX_LOGO_BYTES + 1)
        elif parsed.scheme in ('', 'file'):
            if parsed.scheme == 'file' and parsed.netloc not in ('', 'localhost'):
                raise ValueError('Use a local file path or HTTP(S) URL for Company logo.')
            path = Path(unquote(parsed.path) if parsed.scheme else source)
            with path.open('rb') as stream:
                raw = stream.read(MAX_LOGO_BYTES + 1)
        else:
            raise ValueError('Company logo must be a file path or HTTP(S) URL.')
        if len(raw) > MAX_LOGO_BYTES:
            raise ValueError('Company logo must be at most 10 MB.')
        with Image.open(io.BytesIO(raw)) as image:
            if image.format not in ('PNG', 'JPEG'):
                raise ValueError('Company logo must be a PNG or JPEG image.')
            if image.width * image.height > 20_000_000:
                raise ValueError('Company logo has too many pixels (maximum 20 million).')
            image.load()
            output = io.BytesIO()
            image.convert('RGBA').save(output, format='PNG')
            width, height = image.size
        parts = [document.part]
        for section in document.sections:
            parts.extend((section.header.part, section.footer.part))
        replaced = False
        for part in dict.fromkeys(parts):
            # Replace the existing drawing without changing its position.
            blips = part.element.xpath('.//a:blip')
            for blip in blips:
                rid, _ = part.get_or_add_image(io.BytesIO(output.getvalue()))
                blip.set(qn('r:embed'), rid)
                blip.attrib.pop(qn('r:link'), None)
                drawing = next((node for node in blip.iterancestors() if node.tag == qn('w:drawing')), None)
                if drawing is not None:
                    extents = drawing.xpath('.//wp:extent')
                    if extents:
                        maxw, maxh = int(extents[0].get('cx')), int(extents[0].get('cy'))
                        scale = min(maxw / width, maxh / height)
                        for extent in extents + drawing.xpath('.//a:xfrm/a:ext'):
                            extent.set('cx', str(int(width * scale)))
                            extent.set('cy', str(int(height * scale)))
                replaced = True
        if not replaced:
            # Match the original template's floating logo at the upper left.
            # Anchoring keeps the business text and invoice tables in place.
            scale = min(2325375 / width, 2200940 / height)
            shape = document.paragraphs[0].add_run().add_picture(
                io.BytesIO(output.getvalue()), width=int(width * scale), height=int(height * scale))
            inline = shape._inline
            anchor = OxmlElement('wp:anchor')
            for key, value in {'distT':'0','distB':'0','distL':'114300','distR':'114300',
                               'simplePos':'0','relativeHeight':'251658240','behindDoc':'0',
                               'locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():
                anchor.set(key, value)
            simple = OxmlElement('wp:simplePos'); simple.set('x','0'); simple.set('y','0'); anchor.append(simple)
            for axis, relative, offset in [('H','column','-478466'),('V','paragraph','-457200')]:
                position = OxmlElement('wp:position'+axis); position.set('relativeFrom',relative)
                value = OxmlElement('wp:posOffset'); value.text=offset; position.append(value); anchor.append(position)
            anchor.append(inline.find(qn('wp:extent')))
            anchor.append(OxmlElement('wp:wrapNone'))
            for tag in ('wp:docPr','wp:cNvGraphicFramePr','a:graphic'):
                anchor.append(inline.find(qn(tag)))
            inline.getparent().replace(inline, anchor)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Unable to load Company logo. Use a readable PNG/JPEG path inside the container or an accessible HTTP(S) URL.') from error

def generate(data, settings, number):
    private_template = Path(os.getenv('DATA_DIR','/data')) / 'invoice-template.docx'
    d = Document(private_template if private_template.exists() else BASE / 'templates/invoice.docx')
    # Paragraph positions match both the original and sanitized template.
    for i,key in enumerate(('name','business','address','phone','email','website')):
        p=d.paragraphs[i]
        nodes=p._p.xpath('.//w:t')
        if nodes:
            nodes[0].text=settings[key]
            for n in nodes[1:]: n.text=''
        else: text(p,settings[key])
    recipient = 'Invoice to: ' + data['customer']
    address = data.get('customer_address', '').strip()
    if address:
        recipient += '\n' + address
    text(d.paragraphs[7], recipient)
    text(d.paragraphs[12], 'Date of invoice: ' + uk(data['date']))
    text(d.paragraphs[13], 'Payment terms: ' + settings['terms'])
    for i,key in ((4,'email'),(5,'website')):
        for link in d.paragraphs[i]._p.xpath('.//w:hyperlink'):
            rid=link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            if rid in d.part.rels:
                value=settings[key]
                d.part.rels[rid]._target=('mailto:'+value if key=='email' else (value if value.startswith(('http://','https://')) else 'https://'+value)) if value else ''
    business = os.getenv('BUSINESS_NAME', '').strip()
    if business:
        text(d.paragraphs[1], business)
    company_logo(d)
    # ISO dates sort chronologically; Python's stable sort retains same-day order.
    lines = sorted(data['lines'], key=lambda line: date.fromisoformat(line['date']))
    table = d.tables[0]
    # Clone a formatted service row, then replace the six template placeholders.
    prototype = deepcopy(table.rows[2]._tr)
    deposit_row, total_row = table.rows[8], table.rows[9]
    for row in list(table.rows)[2:8]:
        table._tbl.remove(row._tr)
    for line in lines:
        row_xml = deepcopy(prototype)
        deposit_row._tr.addprevious(row_xml)
        row = _Row(row_xml, table)
        properties = row_xml.get_or_add_trPr()
        if properties.find(qn('w:cantSplit')) is None:
            properties.append(OxmlElement('w:cantSplit'))
        values = [uk(line['date']), line['venue'], line['service'], money(line['fee'])]
        for c, v in zip(row.cells, values):
            cell(c, v)
    for row in list(table.rows)[:2]:
        properties = row._tr.get_or_add_trPr()
        if properties.find(qn('w:tblHeader')) is None:
            properties.append(OxmlElement('w:tblHeader'))
    cell(deposit_row.cells[-1], money(data['deposit']))
    cell(total_row.cells[-1], money(sum(x['fee'] for x in data['lines'])-data['deposit']))
    cell(d.tables[1].rows[0].cells[1], number)
    bank = d.tables[2]
    for row, label, value in [(1,'Account Name',settings['account_name']), (2,'Account Number',settings['account_number']), (3,'Sort Code',settings['sort_code'])]:
        cell(bank.rows[row].cells[0],label); cell(bank.rows[row].cells[1],value)
    with tempfile.TemporaryDirectory() as temp:
        root=Path(temp); source=root/'invoice.docx'; d.save(source)
        command = [os.getenv('SOFFICE','soffice'), f'-env:UserInstallation={ (root/"profile").as_uri() }', '--headless', '--convert-to', 'pdf', '--outdir', str(root), str(source)]
        result=subprocess.run(command,capture_output=True,timeout=90)
        pdf=root/'invoice.pdf'
        if result.returncode or not pdf.exists(): raise RuntimeError('PDF conversion failed. Check the container logs.')
        return pdf.read_bytes()
